"""LocalTranscriber GPU Edition: offline transcription with optional diarization.

Run: .venv\\Scripts\\python.exe gui_transcriber.py
Whisper and pyannote run sequentially in separate Python processes. Audio stays local.
"""
from __future__ import annotations

import os
import queue
import subprocess
import tempfile
import threading
import sys
import json
import time
import traceback
from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, scrolledtext, simpledialog, ttk

APP_DIR = Path(__file__).resolve().parent
OUT_DIR = APP_DIR / "output"
DIAR_PYTHON = APP_DIR / ".venv-diar" / "Scripts" / "python.exe"
DIAR_SCRIPT = APP_DIR / "diarize.py"
MERGE_SCRIPT = APP_DIR / "merge_transcript.py"
WHISPER_WORKER = APP_DIR / "whisper_worker.py"
FFMPEG7_BIN = APP_DIR / "tools" / "ffmpeg7" / "ffmpeg-7.1.1-full_build-shared" / "bin"

MEDIA_TYPES = (
    ("Аудио и видео", "*.mkv *.mp4 *.mp3 *.wav *.m4a *.mov *.aac *.flac *.ogg *.webm *.wma"),
    ("Все файлы", "*.*"),
)
DEFAULT_TERMS = (
    "YouGile, CRM-система, 1С:УНФ, Моё дело, управленческий учёт, "
    "номенклатура, коммерческое предложение, КП, "
    "рентабельность, маржинальность, дебиторская задолженность"
)



def ffmpeg_executable() -> str:
    """Prefer validated FFmpeg 7 Shared, required by this GPU workstation's TorchCodec."""
    ffmpeg7 = FFMPEG7_BIN / ("ffmpeg.exe" if os.name == "nt" else "ffmpeg")
    if ffmpeg7.is_file():
        return str(ffmpeg7)
    # CPU-only installations may use a separately installed compatible FFmpeg.
    import shutil
    return shutil.which("ffmpeg") or ""


def clock(seconds: float) -> str:
    seconds = max(0, int(seconds))
    h, rem = divmod(seconds, 3600)
    m, s = divmod(rem, 60)
    return f"{h:02d}:{m:02d}:{s:02d}"


def srt_time(seconds: float) -> str:
    milliseconds = max(0, round(float(seconds) * 1000))
    h, rem = divmod(milliseconds, 3_600_000)
    m, rem = divmod(rem, 60_000)
    s, ms = divmod(rem, 1000)
    return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"


def get_unique_stem(source: Path, model: str) -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    base = OUT_DIR / f"{source.stem}_{model}_{stamp}"
    result = base
    suffix = 2
    while any(Path(str(result) + ext).exists() for ext in (".txt", ".srt", ".docx")):
        result = OUT_DIR / f"{base.name}_{suffix}"
        suffix += 1
    return result


def export_files(stem: Path, segments: list[tuple[float, float, str]], source: Path) -> list[Path]:
    """Preserve the pre-existing raw Whisper TXT/SRT/DOCX outputs."""
    txt = Path(str(stem) + ".txt")
    srt = Path(str(stem) + ".srt")
    txt_lines = [f"[{clock(start)} – {clock(end)}] {text}" for start, end, text in segments]
    txt.write_text("\n".join(txt_lines) + "\n", encoding="utf-8-sig")
    srt_lines = []
    for i, (start, end, text) in enumerate(segments, 1):
        srt_lines.append(f"{i}\n{srt_time(start)} --> {srt_time(end)}\n{text}\n")
    srt.write_text("\n".join(srt_lines), encoding="utf-8-sig")

    paths = [txt, srt]
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError(
            "Для экспорта Word нужен python-docx. "
            "Выполните: .\\.venv\\Scripts\\python.exe -m pip install python-docx. "
            f"TXT и SRT уже сохранены: {txt}, {srt}"
        ) from exc

    doc = Document()
    doc.add_heading("Транскрипция встречи", level=1)
    doc.add_paragraph(f"Исходный файл: {source.name}")
    for start, end, text in segments:
        doc.add_paragraph(f"[{clock(start)} – {clock(end)}] {text}")
    docx = Path(str(stem) + ".docx")
    doc.save(docx)
    paths.append(docx)
    return paths


def export_speaker_docx(stem: Path, merged: list[dict], source: Path) -> Path:
    from docx import Document

    doc = Document()
    doc.add_heading("Транскрипция встречи — по спикерам", 0)
    doc.add_paragraph(f"Исходный файл: {source.name}")
    doc.add_paragraph(
        "Метки говорящих определены автоматически и требуют проверки по аудио. "
        "Если один сегмент содержит несколько голосов, разметка может быть неточной."
    )
    for item in merged:
        p = doc.add_paragraph()
        heading = p.add_run(
            f"[{clock(item['start'])} – {clock(item['end'])}] {item['speaker']}: "
        )
        heading.bold = True
        p.add_run(item["text"])
    path = Path(str(stem) + ".docx")
    doc.save(path)
    return path


def run_local_command(command: list[str], *, env: dict[str, str] | None = None, secret: str = "") -> str:
    """Do not pass secret values via command-line args or persist subprocess logs."""
    result = subprocess.run(
        command,
        cwd=str(APP_DIR),
        env=env,
        stdin=subprocess.DEVNULL,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        check=False,
    )
    if result.returncode:
        output = (result.stderr or "") + "\n" + (result.stdout or "")
        if secret:
            output = output.replace(secret, "[токен скрыт]")
        raise RuntimeError(
            "Команда завершилась с ошибкой. Код: " + str(result.returncode) + "\n" + output[-2500:]
        )
    return result.stdout


class TranscriberApp:
    def __init__(self, root: tk.Tk):
        self.root = root
        self.root.title("LocalTranscriber GPU Edition — локальная транскрибация")
        self.root.geometry("920x775")
        self.root.minsize(720, 630)
        self.jobs: queue.Queue = queue.Queue()
        self.busy = False
        self.latest_outputs: list[Path] = []

        self.file_path = tk.StringVar()
        self.model_value = tk.StringVar(value="small")
        self.whisper_device = tk.StringVar(value="GPU (CUDA)")
        self.diar_device = tk.StringVar(value="GPU (CUDA)")
        self.use_terms = tk.BooleanVar(value=True)
        self.use_speakers = tk.BooleanVar(value=False)
        self.speaker_count = tk.StringVar(value="2")
        self.status = tk.StringVar(value="Готово к работе. Укажи файл и нажми «Распознать».")
        self.progress_text = tk.StringVar(value="")
        self._build()
        self.root.after(120, self._poll)
        self.root.protocol("WM_DELETE_WINDOW", self._on_close)

    def _build(self):
        body = ttk.Frame(self.root, padding=18)
        body.pack(fill="both", expand=True)
        ttk.Label(body, text="LocalTranscriber", font=("Segoe UI", 20, "bold")).pack(anchor="w")
        ttk.Label(body, text="Видео и аудио обрабатываются локально, без платных API.").pack(anchor="w", pady=(1, 14))

        source_box = ttk.LabelFrame(body, text="Исходная запись", padding=10)
        source_box.pack(fill="x", pady=(0, 12))
        pathrow = ttk.Frame(source_box)
        pathrow.pack(fill="x")
        ttk.Entry(pathrow, textvariable=self.file_path).pack(side="left", fill="x", expand=True)
        self.browse_button = ttk.Button(pathrow, text="Выбрать файл…", command=self._browse)
        self.browse_button.pack(side="left", padx=(8, 0))
        ttk.Label(source_box, text="Большие MKV/MP4 обрабатываются с диска, без загрузки в браузер.", foreground="#60666d").pack(anchor="w", pady=(6, 0))

        settings = ttk.LabelFrame(body, text="Параметры распознавания", padding=10)
        settings.pack(fill="x", pady=(0, 12))
        row = ttk.Frame(settings)
        row.pack(fill="x")
        ttk.Label(row, text="Модель:").pack(side="left")
        self.model_combo = ttk.Combobox(row, textvariable=self.model_value, values=("base", "small", "medium"), state="readonly", width=12)
        self.model_combo.pack(side="left", padx=(8, 18))
        ttk.Label(row, text="Whisper:").pack(side="left")
        self.whisper_device_combo = ttk.Combobox(
            row, textvariable=self.whisper_device, values=("GPU (CUDA)", "CPU"),
            state="readonly", width=13,
        )
        self.whisper_device_combo.pack(side="left", padx=(7, 14))
        ttk.Label(row, text="Русский • beam=5 • VAD").pack(side="left")
        self.terms_checkbox = ttk.Checkbutton(settings, text="Подсказывать специальные термины модели", variable=self.use_terms)
        self.terms_checkbox.pack(anchor="w", pady=(10, 4))
        self.terms = scrolledtext.ScrolledText(settings, height=3, wrap="word", font=("Segoe UI", 10))
        self.terms.insert("1.0", DEFAULT_TERMS)
        self.terms.pack(fill="x")
        ttk.Label(settings, text="Подсказка повышает шанс распознать термины, но иногда навязывает слова.", foreground="#60666d").pack(anchor="w", pady=(4, 0))

        speaker_box = ttk.LabelFrame(body, text="Разделение говорящих (экспериментально)", padding=10)
        speaker_box.pack(fill="x", pady=(0, 12))
        speaker_row = ttk.Frame(speaker_box)
        speaker_row.pack(fill="x")
        self.speakers_checkbox = ttk.Checkbutton(
            speaker_row,
            text="Определять спикеров (pyannote Community-1)",
            variable=self.use_speakers,
            command=self._toggle_speakers,
        )
        self.speakers_checkbox.pack(side="left")
        ttk.Label(speaker_row, text="Участников:").pack(side="left", padx=(24, 6))
        self.speakers_combo = ttk.Combobox(
            speaker_row, textvariable=self.speaker_count,
            values=("Автоматически", "2", "3", "4"), state="disabled", width=16,
        )
        self.speakers_combo.pack(side="left", padx=(0, 14))
        ttk.Label(speaker_row, text="Устройство:").pack(side="left")
        self.diar_device_combo = ttk.Combobox(
            speaker_row, textvariable=self.diar_device,
            values=("GPU (CUDA)", "CPU"), state="disabled", width=13,
        )
        self.diar_device_combo.pack(side="left", padx=(7, 0))
        ttk.Label(
            speaker_box,
            text="Модели выполняются по очереди в отдельных процессах. "
                 "Для GPU требуется RTX NVIDIA; для pyannote — FFmpeg 7 Shared и Hugging Face.",
            foreground="#60666d", wraplength=850,
        ).pack(anchor="w", pady=(6, 0))

        actions = ttk.Frame(body)
        actions.pack(fill="x", pady=(0, 8))
        self.start_button = ttk.Button(actions, text="▶  Распознать", command=self._start)
        self.start_button.pack(side="left")
        self.folder_button = ttk.Button(actions, text="Открыть результаты", command=self._open_folder)
        self.folder_button.pack(side="left", padx=(8, 0))
        ttk.Label(actions, text="Экспорт: TXT + SRT + Word (DOCX)").pack(side="right")

        ttk.Label(body, textvariable=self.status, wraplength=850).pack(anchor="w", pady=(3, 3))
        self.progress = ttk.Progressbar(body, mode="indeterminate", length=450)
        self.progress.pack(fill="x", pady=(4, 4))
        ttk.Label(body, textvariable=self.progress_text, foreground="#60666d", wraplength=850).pack(anchor="w", pady=(0, 10))

        result_frame = ttk.LabelFrame(body, text="Распознанный текст (предпросмотр)", padding=8)
        result_frame.pack(fill="both", expand=True)
        self.preview = scrolledtext.ScrolledText(result_frame, wrap="word", font=("Segoe UI", 10), state="disabled")
        self.preview.pack(fill="both", expand=True)

    def _toggle_speakers(self):
        if not self.busy:
            state = "readonly" if self.use_speakers.get() else "disabled"
            self.speakers_combo.configure(state=state)
            self.diar_device_combo.configure(state=state)

    def _browse(self):
        selected = filedialog.askopenfilename(title="Выберите видео или аудиофайл", filetypes=MEDIA_TYPES)
        if selected:
            self.file_path.set(selected)

    def _set_preview(self, text: str):
        self.preview.configure(state="normal")
        self.preview.delete("1.0", "end")
        self.preview.insert("1.0", text)
        self.preview.configure(state="disabled")

    def _start(self):
        if self.busy:
            return
        path = Path(self.file_path.get().strip().strip('"'))
        if not path.is_file():
            messagebox.showwarning("Не выбран файл", "Выбери существующий файл с аудио или видео.")
            return
        if not WHISPER_WORKER.is_file():
            messagebox.showerror("Отсутствует модуль Whisper", "Нужен файл whisper_worker.py из GPU-обновления.")
            return
        model_name = self.model_value.get()
        whisper_device = "cuda" if self.whisper_device.get() == "GPU (CUDA)" else "cpu"
        diar_device = "cuda" if self.diar_device.get() == "GPU (CUDA)" else "cpu"
        glossary = self.terms.get("1.0", "end").strip() if self.use_terms.get() else ""
        with_speakers = self.use_speakers.get()
        count = 0 if self.speaker_count.get() == "Автоматически" else int(self.speaker_count.get())
        token = ""
        if with_speakers:
            if not DIAR_PYTHON.is_file() or not DIAR_SCRIPT.is_file() or not MERGE_SCRIPT.is_file():
                messagebox.showerror(
                    "Нет модулей разделения спикеров",
                    "Проверь, что файлы diarize.py, merge_transcript.py и "
                    ".venv-diar\\Scripts\\python.exe находятся в папке LocalTranscriber.",
                )
                return
            if not ffmpeg_executable():
                messagebox.showerror("FFmpeg не найден", "Не найден ffmpeg.exe. Проверь tools/ffmpeg7 или PATH.")
                return
            if not FFMPEG7_BIN.is_dir() and diar_device == "cuda":
                messagebox.showerror("FFmpeg 7 не найден", "Для TorchCodec нужен tools/ffmpeg7/ffmpeg-7.1.1-full_build-shared/bin.")
                return
            entered = simpledialog.askstring(
                "Hugging Face — локальная модель",
                "Введите Hugging Face read token. Он не будет сохранён в файлы.\n"
                "Если модель уже доступна через сохранённый вход Hugging Face, "
                "можно оставить поле пустым.",
                show="*", parent=self.root,
            )
            if entered is None:
                return
            token = entered.strip()

        self.busy = True
        self.start_button.configure(state="disabled")
        self.browse_button.configure(state="disabled")
        self.model_combo.configure(state="disabled")
        self.whisper_device_combo.configure(state="disabled")
        self.diar_device_combo.configure(state="disabled")
        self.terms_checkbox.configure(state="disabled")
        self.terms.configure(state="disabled")
        self.speakers_checkbox.configure(state="disabled")
        self.speakers_combo.configure(state="disabled")
        self.status.set("Запуск распознавания…")
        self.progress_text.set("Распознавание и определение спикеров выполняются последовательно.")
        self.progress.configure(mode="indeterminate")
        self.progress.start(12)
        self._set_preview("")
        worker = threading.Thread(
            target=self._transcribe,
            args=(path, model_name, glossary, with_speakers, count, token, whisper_device, diar_device),
            daemon=True,
        )
        worker.start()

    def _transcribe(
        self, path: Path, model_name: str, glossary: str, with_speakers: bool,
        count: int, token: str, whisper_device: str, diar_device: str,
    ):
        raw_paths: list[Path] = []
        start_wall = time.perf_counter()
        try:
            OUT_DIR.mkdir(parents=True, exist_ok=True)
            # Whisper lives in a separate process: no persistent CUDA context remains
            # in the GUI process when the pyannote subprocess starts on a 4GB GPU.
            with tempfile.TemporaryDirectory(prefix="lt_whisper_", dir=str(OUT_DIR)) as scratch:
                scratch_dir = Path(scratch)
                whisper_json = scratch_dir / "whisper.json"
                glossary_file = scratch_dir / "glossary.txt"
                if glossary:
                    glossary_file.write_text(glossary, encoding="utf-8")
                command = [
                    sys.executable, str(WHISPER_WORKER),
                    "--input", str(path), "--model", model_name,
                    "--device", whisper_device, "--output", str(whisper_json),
                ]
                if glossary:
                    command += ["--glossary-file", str(glossary_file)]
                self.jobs.put(("status", f"Whisper {model_name} ({whisper_device.upper()}): распознаю запись…"))
                run_local_command(command)
                payload = json.loads(whisper_json.read_text(encoding="utf-8"))
                whisper_load = float(payload.get("model_load_seconds", 0))
                whisper_run = float(payload.get("processing_seconds", 0))
                self.jobs.put(("position", f"Whisper: загрузка {whisper_load:.1f} с, распознавание {whisper_run:.1f} с"))
                gathered = [
                    (float(s["start"]), float(s["end"]), str(s["text"]))
                    for s in payload["segments"]
                ]
            if not gathered:
                raise RuntimeError("Речь не обнаружена. Проверь звуковую дорожку записи.")
            stem = get_unique_stem(path, model_name)
            raw_paths = export_files(stem, gathered, path)
            elapsed = time.perf_counter() - start_wall
            raw_preview = "\n".join(f"[{clock(a)} – {clock(b)}] {c}" for a, b, c in gathered)
            if not with_speakers:
                self.jobs.put(("done", raw_paths, raw_preview, len(gathered), elapsed, f"Whisper: {whisper_run:.1f} с; загрузка: {whisper_load:.1f} с."))
                return

            self.jobs.put(("status", "Подготавливаю аудио 16 кГц для определения спикеров…"))
            with tempfile.TemporaryDirectory(prefix="lt_diar_", dir=str(OUT_DIR)) as scratch:
                wav = Path(scratch) / "mono_16khz.wav"
                run_local_command([
                    ffmpeg_executable(), "-hide_banner", "-loglevel", "error", "-nostdin", "-y",
                    "-i", str(path), "-vn", "-ac", "1", "-ar", "16000",
                    "-c:a", "pcm_s16le", str(wav),
                ])
                diar_prefix = Path(str(stem) + "_diarization")
                env = os.environ.copy()
                env["PYANNOTE_METRICS_ENABLED"] = "0"
                env["HF_HUB_DISABLE_TELEMETRY"] = "1"
                if FFMPEG7_BIN.is_dir():
                    env["FFMPEG_DLL_DIR"] = str(FFMPEG7_BIN)
                    env["PATH"] = str(FFMPEG7_BIN) + os.pathsep + env.get("PATH", "")
                if token:
                    env["HF_TOKEN"] = token
                self.jobs.put(("status", f"pyannote ({diar_device.upper()}): определяю говорящих…"))
                self.jobs.put(("position", "Whisper завершён, модель выгружена; выполняется pyannote."))
                try:
                    run_local_command([
                        str(DIAR_PYTHON), str(DIAR_SCRIPT),
                        "--input", str(wav),
                        "--num-speakers", str(count),
                        "--device", diar_device,
                        "--output-prefix", str(diar_prefix),
                        "--non-interactive",
                    ], env=env, secret=token)
                finally:
                    env.pop("HF_TOKEN", None)
                    token = ""

            from merge_transcript import Subtitle, read_diarization, merge, save_files
            self.jobs.put(("status", "Объединяю текст Whisper с метками говорящих…"))
            turns, diar_meta, source_key = read_diarization(Path(str(diar_prefix) + ".json"))
            subtitles = [Subtitle(a, b, text) for a, b, text in gathered]
            merged = merge(subtitles, turns, threshold=0.60)
            speaker_prefix = Path(str(stem) + "_speakers")
            metadata = {
                "source_audio": path.name,
                "source_diarization": Path(str(diar_prefix) + ".json").name,
                "diarization_field": source_key,
                "whisper_device": whisper_device,
                "pyannote_device": diar_device,
                "method": "Maximum voiced time overlap per Whisper segment",
                "warning": "Speaker changes inside a Whisper segment may require word-level alignment; review manually.",
            }
            speaker_paths = save_files(speaker_prefix, merged, metadata)
            speaker_paths.append(export_speaker_docx(speaker_prefix, merged, path))
            interval_paths = [Path(str(diar_prefix) + ".json"), Path(str(diar_prefix) + ".txt")]
            all_paths = raw_paths + speaker_paths + interval_paths
            preview = "\n".join(
                f"[{clock(item['start'])} – {clock(item['end'])}] {item['speaker']}: {item['text']}"
                for item in merged
            )
            assigned = sum(item["status"] == "assigned" for item in merged)
            pyannote_load = float(diar_meta.get("model_load_seconds", 0))
            pyannote_run = float(diar_meta.get("processing_seconds", 0))
            note = (
                f"Whisper: {whisper_run:.1f} с; pyannote: {pyannote_run:.1f} с "
                f"(загрузка модели {pyannote_load:.1f} с). "
                f"Спикеры: {assigned} из {len(merged)} фрагментов назначены. Проверь по аудио."
            )
            self.jobs.put(("done", all_paths, preview, len(merged), time.perf_counter() - start_wall, note))
        except Exception:
            error_text = traceback.format_exc()
            if token:
                error_text = error_text.replace(token, "[токен скрыт]")
            if raw_paths:
                self.jobs.put(("partial", raw_paths, error_text))
            else:
                self.jobs.put(("error", error_text))

    def _poll(self):
        try:
            while True:
                message = self.jobs.get_nowait()
                kind = message[0]
                if kind == "status":
                    self.status.set(message[1])
                elif kind == "position":
                    self.progress_text.set(message[1])
                elif kind == "error":
                    self._finish()
                    self.status.set("Ошибка обработки. Подробности показаны ниже.")
                    self.progress_text.set("Проверьте установку библиотек и доступ к моделям.")
                    self._set_preview(message[1])
                    messagebox.showerror("Ошибка", "Не удалось обработать файл. Подробности показаны в окне программы.")
                elif kind == "partial":
                    _, paths, error = message
                    self.latest_outputs = paths
                    self._finish()
                    self.status.set("Whisper завершён, но определение спикеров не удалось. Основная транскрипция сохранена.")
                    self.progress_text.set("Файлы: " + ", ".join(p.name for p in paths))
                    self._set_preview(error)
                    messagebox.showwarning("Частичный результат", "Текст Whisper сохранён. Ошибка pyannote показана в окне программы.")
                elif kind == "done":
                    _, paths, transcript, count, elapsed, note = message
                    self.latest_outputs = paths
                    self._finish()
                    self.status.set(f"Готово! {count} фрагментов, время обработки {elapsed:.1f} с.")
                    main_files = [p.name for p in paths if "_speakers" in p.stem] or [p.name for p in paths]
                    self.progress_text.set((note + "  " if note else "") + "Сохранено: " + ", ".join(main_files))
                    self._set_preview(transcript)
        except queue.Empty:
            pass
        self.root.after(120, self._poll)

    def _finish(self):
        self.busy = False
        self.progress.stop()
        self.progress.configure(value=0)
        self.start_button.configure(state="normal")
        self.browse_button.configure(state="normal")
        self.model_combo.configure(state="readonly")
        self.whisper_device_combo.configure(state="readonly")
        self.terms_checkbox.configure(state="normal")
        self.terms.configure(state="normal")
        self.speakers_checkbox.configure(state="normal")
        self._toggle_speakers()

    def _open_folder(self):
        OUT_DIR.mkdir(parents=True, exist_ok=True)
        if os.name == "nt":
            os.startfile(str(OUT_DIR))
        else:
            messagebox.showinfo("Папка с результатами", str(OUT_DIR))

    def _on_close(self):
        if self.busy and not messagebox.askyesno(
            "Прервать работу?",
            "Идёт транскрибация. Если закрыть программу, текущий результат может не сохраниться. Закрыть?",
        ):
            return
        self.root.destroy()


def main():
    root = tk.Tk()
    try:
        ttk.Style().theme_use("clam")
    except tk.TclError:
        pass
    TranscriberApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
